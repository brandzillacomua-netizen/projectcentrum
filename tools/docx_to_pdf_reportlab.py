from pathlib import Path
import argparse
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(r"A:\centrum")
parser = argparse.ArgumentParser()
parser.add_argument("input", nargs="?", default=str(ROOT / "output" / "documents" / "Vyrobnychyi_tsykl_vid_A_do_Ya_SOP.docx"))
parser.add_argument("output", nargs="?", default=str(ROOT / "output" / "pdf" / "Vyrobnychyi_tsykl_vid_A_do_Ya_SOP.pdf"))
args = parser.parse_args()
DOCX = Path(args.input)
OUT = Path(args.output)
OUT.parent.mkdir(parents=True, exist_ok=True)

pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Bold", r"C:\Windows\Fonts\arialbd.ttf"))

NAVY = colors.HexColor("#17365D")
BLUE = colors.HexColor("#DCE6F1")
PALE = colors.HexColor("#F5F7FA")
GRAY = colors.HexColor("#5B6573")
LINE = colors.HexColor("#D9D9D9")


class SOPDocument(BaseDocTemplate):
    def __init__(self, filename):
        super().__init__(
            filename,
            pagesize=A4,
            leftMargin=21 * mm,
            rightMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
            title="CENTRUM MES Виробничий цикл від А до Я",
            author="Операційний департамент",
            subject="Корпоративний регламент SOP",
        )
        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="body",
            leftPadding=0,
            rightPadding=0,
            topPadding=0,
            bottomPadding=4 * mm,
        )
        self.addPageTemplates(PageTemplate(id="main", frames=[frame], onPage=self._footer))

    @staticmethod
    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(LINE)
        canvas.setLineWidth(0.35)
        canvas.line(21 * mm, 13 * mm, A4[0] - 18 * mm, 13 * mm)
        canvas.setFont("Arial", 8)
        canvas.setFillColor(GRAY)
        canvas.drawCentredString(
            A4[0] / 2,
            8.5 * mm,
            f"CENTRUM MES  |  Виробничий цикл  |  Версія 1.0  |  Сторінка {doc.page}",
        )
        canvas.restoreState()


styles = getSampleStyleSheet()
BODY = ParagraphStyle(
    "BodyUA",
    parent=styles["BodyText"],
    fontName="Arial",
    fontSize=9.0,
    leading=11.5,
    textColor=colors.black,
    spaceAfter=4.5,
)
TITLE = ParagraphStyle(
    "TitleUA",
    parent=BODY,
    fontName="Arial-Bold",
    fontSize=25,
    leading=29,
    spaceAfter=13,
)
SUBTITLE = ParagraphStyle(
    "SubtitleUA",
    parent=BODY,
    fontSize=14,
    leading=18,
    textColor=GRAY,
    spaceAfter=12,
)
H1 = ParagraphStyle(
    "H1UA",
    parent=BODY,
    fontName="Arial-Bold",
    fontSize=16,
    leading=20,
    textColor=NAVY,
    spaceBefore=9,
    spaceAfter=7,
    keepWithNext=True,
)
H2 = ParagraphStyle(
    "H2UA",
    parent=BODY,
    fontName="Arial-Bold",
    fontSize=11.5,
    leading=15,
    textColor=NAVY,
    spaceBefore=7,
    spaceAfter=4,
    keepWithNext=True,
)
H3 = ParagraphStyle(
    "H3UA",
    parent=BODY,
    fontName="Arial-Bold",
    fontSize=10.2,
    leading=13.5,
    spaceBefore=5,
    spaceAfter=3,
    keepWithNext=True,
)
BULLET = ParagraphStyle(
    "BulletUA",
    parent=BODY,
    leftIndent=7 * mm,
    firstLineIndent=-4 * mm,
    bulletIndent=2.5 * mm,
    spaceAfter=3,
)
SMALL = ParagraphStyle(
    "SmallUA",
    parent=BODY,
    fontSize=8.2,
    leading=10.5,
    textColor=GRAY,
)
CELL = ParagraphStyle("CellUA", parent=BODY, fontSize=8.3, leading=10.5, spaceAfter=0)
CELL_BOLD = ParagraphStyle(
    "CellBoldUA", parent=CELL, fontName="Arial-Bold", textColor=NAVY
)
CELL_HEADER = ParagraphStyle(
    "CellHeaderUA", parent=CELL, fontName="Arial-Bold", textColor=colors.white
)


def paragraph_markup(paragraph):
    chunks = []
    for run in paragraph.runs:
        text = escape(run.text).replace("\n", "<br/>")
        if not text:
            continue
        if run.bold:
            text = f"<b>{text}</b>"
        if run.italic:
            text = f"<i>{text}</i>"
        chunks.append(text)
    return "".join(chunks) or escape(paragraph.text)


def has_page_break(paragraph):
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def paragraph_flowable(paragraph):
    text = paragraph_markup(paragraph).strip()
    if not text:
        return None
    name = paragraph.style.name if paragraph.style else ""
    if name == "Title":
        style = TITLE
    elif name == "Subtitle":
        style = SUBTITLE
    elif name == "Heading 1":
        style = H1
    elif name == "Heading 2":
        style = H2
    elif name == "Heading 3":
        style = H3
    elif name == "SOP Small":
        style = SMALL
    elif name.startswith("List Bullet"):
        level = 1 if name.endswith("2") else 0
        style = ParagraphStyle(
            f"Bullet{level}-{id(paragraph)}",
            parent=BULLET,
            leftIndent=(7 + level * 5) * mm,
        )
        text = "• " + text
    else:
        style = BODY
    return Paragraph(text, style)


def table_flowable(table, table_index):
    rows = []
    for r_i, row in enumerate(table.rows):
        values = []
        for c_i, cell in enumerate(row.cells):
            cell_text = "<br/>".join(
                paragraph_markup(p).strip() for p in cell.paragraphs if paragraph_markup(p).strip()
            )
            if r_i == 0 and table_index > 0:
                style = CELL_HEADER
            elif r_i == 0 or (table_index == 0 and c_i == 0):
                style = CELL_BOLD
            else:
                style = CELL
            values.append(Paragraph(cell_text or " ", style))
        rows.append(values)

    col_count = max(len(r) for r in rows)
    usable = A4[0] - 39 * mm
    if col_count == 2:
        widths = [usable * 0.31, usable * 0.69]
    elif col_count == 3:
        widths = [usable * 0.25, usable * 0.36, usable * 0.39]
    elif col_count == 4:
        widths = [usable * 0.18, usable * 0.32, usable * 0.23, usable * 0.27]
    else:
        widths = [usable / col_count] * col_count

    result = Table(rows, colWidths=widths, repeatRows=1 if table_index > 0 else 0, hAlign="LEFT")
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.45, LINE),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if table_index == 0:
        commands.append(("BACKGROUND", (0, 0), (0, -1), BLUE))
    else:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ]
        )
    for r_i in range(1 if table_index > 0 else 0, len(rows)):
        if r_i % 2 == 0:
            commands.append(("BACKGROUND", (0, r_i), (-1, r_i), PALE))
    result.setStyle(TableStyle(commands))
    return result


source = Document(DOCX)
story = []
table_index = 0
for child in source.element.body.iterchildren():
    if child.tag == qn("w:p"):
        paragraph = next(p for p in source.paragraphs if p._p is child)
        if has_page_break(paragraph):
            story.append(PageBreak())
        flowable = paragraph_flowable(paragraph)
        if flowable is not None:
            style_name = paragraph.style.name if paragraph.style else ""
            if style_name == "Heading 1":
                story.append(Spacer(1, 3.2 * mm))
            elif style_name in ("Heading 2", "Heading 3"):
                story.append(Spacer(1, 1.6 * mm))
            story.append(flowable)
    elif child.tag == qn("w:tbl"):
        table = next(t for t in source.tables if t._tbl is child)
        story.append(Spacer(1, 2 * mm))
        story.append(table_flowable(table, table_index))
        story.append(Spacer(1, 3 * mm))
        table_index += 1

SOPDocument(str(OUT)).build(story)
print(OUT)
